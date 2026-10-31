"""
Unit tests for new two-Excel comparison and flow logic.
"""

import json
import os
import unittest
import pandas as pd
from app import execute_compare_two_excels, compare_two_excels_venn, extract_langflow_response

class TestTwoExcelComparison(unittest.TestCase):
    def setUp(self):
        # Create two simple test DataFrames
        self.df1 = pd.DataFrame({
            "Uygunsuzluk": ["Hata 1", "Hata 2", "Hata 3", "Ortak Hata A", "Ortak Hata B"]
        })
        self.df2 = pd.DataFrame({
            "Uygunsuzluk": ["Hata 4", "Hata 5", "Ortak Hata A", "Ortak Hata B"]
        })

    def test_compare_dataframes_venn(self):
        res = execute_compare_two_excels(
            file_1=self.df1,
            file_2=self.df2,
            file_1_name="Versiyon_1.xlsx",
            file_2_name="Versiyon_2.xlsx"
        )
        self.assertEqual(res["durum"], "basarili")
        v = res["venn_verileri"]
        
        # Intersection
        self.assertEqual(v["intersection_count"], 2)
        self.assertIn("Ortak Hata A", v["intersection_items"])
        self.assertIn("Ortak Hata B", v["intersection_items"])
        
        # Only in 1
        self.assertEqual(v["only_a_count"], 3)
        self.assertIn("Hata 1", v["only_a_items"])
        
        # Only in 2
        self.assertEqual(v["only_b_count"], 2)
        self.assertIn("Hata 4", v["only_b_items"])
        
        # Total unique
        self.assertEqual(v["total_unique"], 7)
        # Jaccard = 2 / 7 = 0.2857
        self.assertAlmostEqual(v["jaccard_similarity"], 2/7, places=3)

    def test_compare_with_sample_excel(self):
        sample_path = os.path.abspath("ornek_uygunsuzluklar.xlsx")
        res_str = compare_two_excels_venn(sample_path, sample_path)
        data = json.loads(res_str)
        self.assertEqual(data["durum"], "basarili")
        # Comparing file with itself should have 100% intersection
        v = data["venn_verileri"]
        self.assertEqual(v["only_a_count"], 0)
        self.assertEqual(v["only_b_count"], 0)
        self.assertEqual(v["jaccard_similarity"], 1.0)

    def test_calculate_text_similarity(self):
        from app import calculate_text_similarity
        # Identical strings
        self.assertEqual(calculate_text_similarity("Pano montaj hatası", "Pano montaj hatası"), 1.0)
        # Word order difference
        sim_order = calculate_text_similarity("Pano montaj hatası", "montaj hatası pano")
        self.assertGreaterEqual(sim_order, 0.85)
        # Completely different
        sim_diff = calculate_text_similarity("Kablo kesiti hatalı", "Yazılım derleme hatası")
        self.assertLess(sim_diff, 0.40)

    def test_compute_fuzzy_venn(self):
        from app import compute_fuzzy_venn
        list_a = ["Pano içi etiketleme eksik", "Kablo gevşek", "Hidrolik sızdırmazlık arızası"]
        list_b = ["etiketleme eksik pano içi", "Kablo gevşek", "Yazılım derleme hatası"]
        
        # Exact match (threshold 1.0)
        res_exact = compute_fuzzy_venn(list_a, list_b, threshold=1.0)
        self.assertEqual(res_exact["intersection_count"], 1) # Only "Kablo gevşek"
        
        # Fuzzy match (threshold 0.80)
        res_fuzzy = compute_fuzzy_venn(list_a, list_b, threshold=0.80)
        self.assertEqual(res_fuzzy["intersection_count"], 2) # Both "Kablo gevşek" and reordered "etiketleme..."
        self.assertEqual(res_fuzzy["only_a_count"], 1)
        self.assertEqual(res_fuzzy["only_b_count"], 1)

    def test_fetch_aselix_models_fallback(self):
        from app import fetch_aselix_models
        # Even with dummy/unreachable URL, it should return models list gracefully
        res = fetch_aselix_models(base_url="https://api.aselixai.aselsan.com.tr/v1", api_key="dummy_key")
        self.assertIn("modeller", res)
        self.assertGreater(len(res["modeller"]), 0)

    def test_batch_comparison_offline_fallback(self):
        from app import compare_items_with_aselix_batch
        # When no API key is provided, it must safely execute via local fuzzy distance
        batch_a = ["Pano etiketleme eksik", "Kablo kopuk"]
        cands_b = ["etiketleme eksik pano", "Hidrolik sızıntı"]
        res = compare_items_with_aselix_batch(
            base_url="https://api.aselixai.aselsan.com.tr/v1",
            api_key="",
            model="aselix-llm-v1",
            items_a_batch=batch_a,
            candidates_b=cands_b,
            threshold=0.8,
            timeout=50
        )
        self.assertEqual(res["durum"], "basarili")
        self.assertEqual(len(res["matched_pairs"]), 1)
        self.assertEqual(res["matched_pairs"][0]["item_a"], "Pano etiketleme eksik")

    def test_batch_endpoints_with_testclient(self):
        from fastapi.testclient import TestClient
        from app import web_app
        client = TestClient(web_app)
        
        # 1. Batch Chunk endpoint
        chunk_res = client.post("/api/compare/batch-chunk", json={
            "batch_index": 1,
            "total_batches": 1,
            "items_a_batch": ["Kablo gevşek"],
            "candidates_b": ["Kablo gevşek", "Optik hata"],
            "threshold": 0.8,
            "api_key": ""
        })
        self.assertEqual(chunk_res.status_code, 200)
        c_data = chunk_res.json()
        self.assertEqual(c_data["durum"], "basarili")
        self.assertEqual(len(c_data["matched_pairs"]), 1)

        # 2. Batch Finalize endpoint
        finalize_res = client.post("/api/compare/batch-finalize", json={
            "doc1_name": "Test1",
            "doc2_name": "Test2",
            "threshold": 0.8,
            "all_items_a": ["Kablo gevşek", "Hata A"],
            "all_items_b": ["Kablo gevşek", "Hata B"],
            "matched_pairs": c_data["matched_pairs"],
            "need_executive_summary": False
        })
        self.assertEqual(finalize_res.status_code, 200)
        f_data = finalize_res.json()
        self.assertEqual(f_data["durum"], "basarili")
        self.assertEqual(f_data["venn_verileri"]["intersection_count"], 1)
        self.assertEqual(f_data["venn_verileri"]["only_a_count"], 1)
        self.assertEqual(f_data["venn_verileri"]["only_b_count"], 1)


class TestAselixBatchPipeline(unittest.TestCase):
    """AselixAI API'si mock'lanarak batch kıyaslamanın uçtan uca doğrulanması."""

    def _mock_response(self, content):
        from unittest.mock import MagicMock
        r = MagicMock()
        r.status_code = 200
        r.json.return_value = {"choices": [{"message": {"content": content}}]}
        return r

    def test_ai_index_matching(self):
        from unittest.mock import patch
        import app
        a = ["Acil durdurma butonu takılıyor", "Röle soketi kırık"]
        b = ["Optik sensör sapması", "Acil stop butonu mekanik takılma"]
        llm = '```json\n[{"a": 1, "b": 2, "similarity": 93}, {"a": 2, "b": null, "similarity": 0}]\n```'
        with patch.object(app.requests, "post", return_value=self._mock_response(llm)):
            res = app.compare_items_with_aselix_batch("http://x/v1", "key", "m", a, b, threshold=0.8)
        self.assertEqual(res["kaynak"], "aselix_ai")
        self.assertEqual(len(res["matched_pairs"]), 1)
        self.assertEqual(res["matched_pairs"][0]["item_a"], a[0])
        self.assertEqual(res["matched_pairs"][0]["item_b"], b[1])
        self.assertIn("tanilama", res)

    def test_timeout_falls_back_to_local(self):
        from unittest.mock import patch
        import app
        with patch.object(app.requests, "post", side_effect=app.requests.exceptions.Timeout()):
            res = app.compare_items_with_aselix_batch(
                "http://x/v1", "key", "m", ["Kablo gevşek"], ["Kablo gevşek", "Hata"], threshold=0.8, timeout=1)
        self.assertEqual(res["kaynak"], "yerel_fallback")
        self.assertTrue(res["api_hatasi"])
        self.assertEqual(len(res["matched_pairs"]), 1)

    def test_shortlist_limits_prompt_size(self):
        from app import shortlist_candidates_for_batch
        b = [f"Uygunsuzluk kaydı numara {i} farklı açıklama" for i in range(500)] + ["Pano etiketi eksik"]
        sl = shortlist_candidates_for_batch(["Pano etiketleme eksik"], b, top_k=6, max_total=60)
        self.assertLessEqual(len(sl), 60)
        self.assertIn("Pano etiketi eksik", sl)

    def test_full_batch_flow_via_endpoints(self):
        """15'erli batch'lerle 40 satır: chunk -> finalize, AI mock'lu."""
        from unittest.mock import patch
        from fastapi.testclient import TestClient
        import app
        client = TestClient(app.web_app)
        a = [f"Hata kodu {i} test sapması" for i in range(40)]
        b = a[:20] + [f"Yeni hata {i}" for i in range(10)]
        llm = "[]"  # AI eşleşme bulmasa da fuzzy tam eşleşmeleri yakalamalı
        pairs = []
        with patch.object(app.requests, "post", return_value=self._mock_response(llm)):
            for i in range(0, 40, 15):
                r = client.post("/api/compare/batch-chunk", json={
                    "api_key": "k", "items_a_batch": a[i:i + 15], "candidates_b": b,
                    "threshold": 0.95, "batch_index": i // 15 + 1, "total_batches": 3})
                self.assertEqual(r.status_code, 200)
                pairs += r.json()["matched_pairs"]
        f = client.post("/api/compare/batch-finalize", json={
            "threshold": 0.95, "all_items_a": a, "all_items_b": b,
            "matched_pairs": pairs, "need_executive_summary": False}).json()
        v = f["venn_verileri"]
        self.assertEqual(v["intersection_count"], 20)
        self.assertEqual(v["only_a_count"], 20)
        self.assertEqual(v["only_b_count"], 10)

    def test_extract_items_endpoint(self):
        from fastapi.testclient import TestClient
        from app import web_app
        client = TestClient(web_app)
        path = os.path.abspath("ornek_uygunsuzluklar.xlsx")
        with open(path, "rb") as f1, open(path, "rb") as f2:
            r = client.post("/api/compare/extract-items",
                            files={"file_1": ("a.xlsx", f1.read()), "file_2": ("b.xlsx", f2.read())},
                            data={"col_1": "auto", "col_2": "auto"})
        d = r.json()
        self.assertEqual(d["durum"], "basarili")
        self.assertGreater(len(d["items_1"]), 0)


class TestKokpitIdentityAndEmbedding(unittest.TestCase):
    """KIMLIK_AKTARIMI_ENTEGRASYON_REHBERI.md uyumlu SSO, Iframe Push, Dış Input ve Embedding testleri."""

    def setUp(self):
        from fastapi.testclient import TestClient
        from app import web_app
        self.client = TestClient(web_app)

    def test_iframe_headers_allow_embedding(self):
        resp = self.client.get("/")
        self.assertEqual(resp.status_code, 200)
        self.assertNotIn("x-frame-options", [k.lower() for k in resp.headers.keys()])
        csp = resp.headers.get("Content-Security-Policy", "")
        self.assertIn("frame-ancestors", csp)
        self.assertIn("http://localhost:3004", csp)

    def test_kokpit_identity_push_and_get(self):
        payload = {
            "type": "kokpit-identity",
            "status": "success",
            "source": "kokpit-iframe-bridge",
            "timestamp": "2026-10-29T10:00:00.000Z",
            "token": "eyJhbGciOiJSUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiIxMjM0NSIsInByZWZlcnJlZF91c2VybmFtZSI6ImFobWV0LnlpbG1heiIsIm5hbWUiOiJBaG1ldCBZxLFsbWF6In0.sig",
            "user": {
                "id": 145,
                "sicil": "12345",
                "username": "ahmet.yilmaz",
                "isim": "Ahmet Yılmaz",
                "mail": "ahmet.yilmaz@sirket.com",
                "picture": "/uploads/photos/12345.jpg",
                "role": "Admin",
                "mudurluk": "UGES Teslimat Yönetim Müdürlüğü",
                "mudurluk_id": 191148,
                "direktorluk": "Program Direktörlüğü",
                "birim": "Planlama Birimi",
                "unvan": "Uzman",
                "hierarchicLevel": 2
            },
            "inputs": {
                "engine_mode": "embedding",
                "threshold": 0.80,
                "items_1": ["Pano etiketleme eksik", "Kablo gevşek"],
                "items_2": ["etiketleme eksik pano", "Sensör hatası"]
            }
        }
        post_res = self.client.post("/api/identity", json=payload)
        self.assertEqual(post_res.status_code, 200)
        data = post_res.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["user"]["sicil"], "12345")
        self.assertEqual(data["user"]["isim"], "Ahmet Yılmaz")
        self.assertTrue(data["token_present"])
        self.assertEqual(data["external_inputs"]["engine_mode"], "embedding")

        get_res = self.client.get("/api/identity")
        self.assertEqual(get_res.status_code, 200)
        gdata = get_res.json()
        self.assertTrue(gdata["authenticated"])
        self.assertEqual(gdata["user"]["mudurluk_id"], 191148)

    def test_invalid_identity_returns_400(self):
        bad_res = self.client.post("/api/identity", json={"user": {}})
        self.assertEqual(bad_res.status_code, 400)
        self.assertIn("error", bad_res.json())

    def test_external_input_with_embeddings_and_auto_compare(self):
        ext_payload = {
            "source": "kokpit-iframe-bridge",
            "items_1": ["Pano içi etiketleme eksik", "Kablo kopuk"],
            "items_2": ["etiketleme eksik pano içi", "Yazılım hatası"],
            "embeddings_1": [[1.0, 0.0, 0.0], [0.0, 1.0, 0.0]],
            "embeddings_2": [[0.98, 0.02, 0.0], [0.0, 0.0, 1.0]],
            "engine_mode": "embedding",
            "threshold": 0.80,
            "auto_compare": True
        }
        res = self.client.post("/api/external-input", json=ext_payload)
        self.assertEqual(res.status_code, 200)
        d = res.json()
        self.assertEqual(d["status"], "success")
        self.assertIsNotNone(d["comparison_result"])
        v = d["comparison_result"]["venn_verileri"]
        self.assertEqual(v["intersection_count"], 1)
        self.assertEqual(v["embedding_source"], "dis_input_hazir_embedding")

    def test_compare_run_with_embedding_engine(self):
        res = self.client.post("/api/compare/run", data={
            "input_mode": "paste",
            "engine_mode": "embedding",
            "threshold": 0.75,
            "text_1": "Pano içi etiketleme eksik\nTopraklama cıvatası gevşek",
            "text_2": "etiketleme eksik pano içi\nOptik sensör hatası"
        })
        self.assertEqual(res.status_code, 200)
        d = res.json()
        self.assertEqual(d["durum"], "basarili")
        self.assertEqual(d["engine_mode"], "embedding")
        self.assertEqual(d["venn_verileri"]["intersection_count"], 1)

    def test_bakiye_siparis_items_endpoint(self):
        res = self.client.get("/api/bakiye-siparis/items")
        self.assertEqual(res.status_code, 200)
        d = res.json()
        self.assertEqual(d["durum"], "basarili")
        self.assertEqual(d["faz_secenekleri"], ["Dry-Run", "FKT", "SKT"])
        self.assertTrue(len(d["kayitlar"]) > 0)
        first = d["kayitlar"][0]
        self.assertIn("proje_adi", first)
        self.assertIn("sd", first)
        self.assertIn("kalem", first)
        self.assertIn("urun_kisa_metni", first)

    def test_compare_run_with_bakiye_meta_and_phases(self):
        meta_1 = json.dumps({
            "proje_adi": "Deneme Proje 1",
            "uges_proje_tanim": "UDeneme1",
            "sd": "987654",
            "kalem": "2",
            "urun_kisa_metni": "Malzeme2",
            "faz": "Dry-Run"
        })
        meta_2 = json.dumps({
            "proje_adi": "Deneme Proje 1",
            "uges_proje_tanim": "UDeneme1",
            "sd": "987654",
            "kalem": "2",
            "urun_kisa_metni": "Malzeme2",
            "faz": "FKT"
        })
        res = self.client.post("/api/compare/run", data={
            "input_mode": "paste",
            "engine_mode": "local",
            "threshold": 0.80,
            "text_1": "Pano etiketleme eksik\nKablo gevşek",
            "text_2": "Pano etiketleme eksik\nSensör arızası",
            "meta_1": meta_1,
            "meta_2": meta_2
        })
        self.assertEqual(res.status_code, 200)
        d = res.json()
        self.assertEqual(d["durum"], "basarili")
        self.assertEqual(d["belge_1_meta"]["faz"], "Dry-Run")
        self.assertEqual(d["belge_2_meta"]["faz"], "FKT")
        self.assertEqual(d["belge_1_meta"]["sd"], "987654")
        self.assertIn("Dry-Run", d["ozet_yorum"])
        self.assertIn("FKT", d["ozet_yorum"])

    def test_save_and_get_dryrun_fkt_skt_kpi(self):
        payload = {
            "proje_adi": "Deneme Proje 1",
            "sd": "987654",
            "kalem": "2",
            "faaliyet_1_turu": "FKT",  # Gönderilse bile backend kural gereği Dry-Run yapmalı
            "faaliyet_2_turu": "SKT",
            "urun_kisa_metni": "Malzeme2",
            "uges_proje_tanim": "UDeneme1",
            "malzeme": "5865-7567-0110",
            "ilk_faaliyet_benzersiz_uygunsuzluklar": [
                "Kablo kesiti standart dışı",
                "Topraklama direnci yüksek",
                "Cıvata tork değeri yetersiz"
            ],
            "ortak_uygunsuzluklar": [
                {
                    "ilk_faaliyet_uygunsuzluk": "Pano içi etiketleme eksik",
                    "ikinci_faaliyet_uygunsuzluk": "Pano içi etiketleme eksik",
                    "benzerlik_yuzdesi": 100.0
                }
            ],
            "ikinci_faaliyet_benzersiz_uygunsuzluklar": [
                "Flanş sızdırmazlığı arızalı",
                "Klemens bağlantısı gevşek",
                "Cıvata tork değeri yetersiz"
            ],
            "reddedilen_eslesmeler": [
                {
                    "ilk_faaliyet_uygunsuzluk": "Cıvata tork değeri yetersiz",
                    "ikinci_faaliyet_uygunsuzluk": "Cıvata tork değeri yetersiz",
                    "eski_benzerlik_yuzdesi": 100.0
                }
            ]
        }
        save_res = self.client.post("/api/bakiye-siparis/save-kpi", json=payload)
        self.assertEqual(save_res.status_code, 200)
        sdata = save_res.json()
        self.assertIn(sdata["durum"], ["basarili", "uyari"])
        self.assertEqual(sdata["kayit"]["faaliyet_1_turu"], "Dry-Run")
        self.assertEqual(sdata["kayit"]["faaliyet_2_turu"], "SKT")
        self.assertEqual(sdata["kayit"]["ilk_faaliyet_benzersiz_sayisi"], 3)
        self.assertEqual(sdata["kayit"]["ortak_uygunsuzluk_sayisi"], 1)
        self.assertEqual(sdata["kayit"]["ikinci_faaliyet_benzersiz_sayisi"], 3)

        list_res = self.client.get("/api/bakiye-siparis/kpi-records")
        self.assertEqual(list_res.status_code, 200)
        ldata = list_res.json()
        self.assertTrue(len(ldata["kayitlar"]) >= 1)

    def test_env_and_db_startup_verification(self):
        status_res = self.client.get("/api/bakiye-siparis/db-status")
        self.assertEqual(status_res.status_code, 200)
        s = status_res.json()
        self.assertTrue(s["env_dosyasi_mevcut"])
        self.assertEqual(s["baglanti"]["dbname"], "ud_kokpit")
        self.assertEqual(s["baglanti"]["schema"], "teslimat_md")
        self.assertIn("kpi_tablosu", s)
        self.assertIn("sunucu", s)
        self.assertIn("app_port", s["sunucu"])
        self.assertEqual(s["sunucu"]["app_port"], 8000)

        # .env üzerinden veya ortam değişkeninden port değiştirilebilirlik testi
        from unittest import mock
        from app import get_server_config
        srv = get_server_config()
        self.assertEqual(srv["port"], 8000)
        with mock.patch.dict(os.environ, {"APP_PORT": "9090"}):
            self.assertEqual(get_server_config()["port"], 9090)
        with mock.patch.dict(os.environ, {"PORT": "7070", "APP_PORT": ""}):
            self.assertEqual(get_server_config()["port"], 7070)

    def test_admin_only_llm_suite_and_role_visibility(self):
        # 1) Standart kullanıcı (User) oturumu açıldığında LLM Suite Flow URL ve API Key uyarısı dönmemeli
        self.client.post(
            "/api/identity",
            json={
                "type": "kokpit-identity",
                "status": "success",
                "source": "kokpit-iframe-bridge",
                "user": {
                    "sicil": "67890",
                    "username": "mehmet.kaya",
                    "isim": "Mehmet Kaya",
                    "role": "User"
                }
            }
        )
        id_res_user = self.client.get("/api/identity").json()
        self.assertFalse(id_res_user["user"]["is_admin"])

        sample_user_res = self.client.post(
            "/api/flow/sample-compare",
            json={
                "base_url": "https://api.aselixai.aselsan.com.tr/v1",
                "flow_id": "c16ecc83-ff64-4fc7-bb52-0f594beaf70a",
                "api_key": ""
            }
        ).json()
        self.assertNotIn("ASELSAN LLM Suite Flow Bağlantısı", sample_user_res.get("flow_analiz_ozeti", ""))
        self.assertNotIn("Hedef Akış:", sample_user_res.get("flow_analiz_ozeti", ""))
        self.assertNotIn("API Key", sample_user_res.get("flow_analiz_ozeti", ""))

        # 2) Admin kullanıcısı oturumu açıldığında LLM Suite Flow URL ve Ayarlar bilgisi görünmeli
        self.client.post(
            "/api/identity",
            json={
                "type": "kokpit-identity",
                "status": "success",
                "source": "kokpit-iframe-bridge",
                "user": {
                    "sicil": "12345",
                    "username": "ahmet.yilmaz",
                    "isim": "Ahmet Yılmaz",
                    "role": "Admin"
                }
            }
        )
        id_res_admin = self.client.get("/api/identity").json()
        self.assertTrue(id_res_admin["user"]["is_admin"])

        sample_admin_res = self.client.post(
            "/api/flow/sample-compare",
            json={
                "base_url": "https://api.aselixai.aselsan.com.tr/v1",
                "flow_id": "c16ecc83-ff64-4fc7-bb52-0f594beaf70a",
                "api_key": ""
            }
        ).json()
        self.assertIn("ASELSAN LLM Suite Flow Bağlantısı", sample_admin_res.get("flow_analiz_ozeti", ""))
        self.assertIn("Hedef Akış:", sample_admin_res.get("flow_analiz_ozeti", ""))
        self.assertIn("Ayarlar", sample_admin_res.get("flow_analiz_ozeti", ""))

        # Temizle
        self.client.delete("/api/identity")

    def test_kpi_records_filtering_and_admin_delete(self):
        # 1) Geçici test KPI kaydı ekle
        test_payload = {
            "proje_adi": "TEST-FILTRE-SIL-PROJESI",
            "sd": "99999999",
            "kalem": "999",
            "faaliyet_1_turu": "Dry-Run",
            "faaliyet_2_turu": "SKT",
            "urun_kisa_metni": "TEST RADAR MODULU",
            "ilk_faaliyet_benzersiz_uygunsuzluklar": ["Test hata A1"],
            "ortak_uygunsuzluklar": [{"ilk_faaliyet_uygunsuzluk": "Ortak X", "ikinci_faaliyet_uygunsuzluk": "Ortak X", "benzerlik_yuzdesi": 100}],
            "ikinci_faaliyet_benzersiz_uygunsuzluklar": ["Test hata B1"]
        }
        save_res = self.client.post("/api/bakiye-siparis/save-kpi", json=test_payload)
        self.assertEqual(save_res.status_code, 200)

        # 2) Tablo filtreleme sorgularını test et
        filter_res = self.client.get(
            "/api/bakiye-siparis/kpi-records",
            params={
                "proje": "TEST-FILTRE-SIL-PROJESI",
                "sd": "99999999",
                "faaliyet_2_turu": "SKT",
                "search": "RADAR"
            }
        )
        self.assertEqual(filter_res.status_code, 200)
        fdata = filter_res.json()
        self.assertEqual(len(fdata["kayitlar"]), 1)
        self.assertEqual(fdata["kayitlar"][0]["urun_kisa_metni"], "TEST RADAR MODULU")

        # 3) Standart kullanıcı (User) silmeye çalışırsa 403 Yetkisiz dönmeli
        self.client.post(
            "/api/identity",
            json={
                "type": "kokpit-identity",
                "status": "success",
                "source": "kokpit-iframe-bridge",
                "user": {"sicil": "67890", "username": "mehmet.kaya", "isim": "Mehmet Kaya", "role": "User"}
            }
        )
        del_unauth = self.client.post("/api/bakiye-siparis/delete-kpi", json=test_payload)
        self.assertEqual(del_unauth.status_code, 403)
        self.assertEqual(del_unauth.json()["durum"], "yetkisiz")

        # 4) Admin kullanıcısı silmeye çalışırsa başarıyla silinmeli
        self.client.post(
            "/api/identity",
            json={
                "type": "kokpit-identity",
                "status": "success",
                "source": "kokpit-iframe-bridge",
                "user": {"sicil": "12345", "username": "ahmet.yilmaz", "isim": "Ahmet Yılmaz", "role": "Admin"}
            }
        )
        del_admin = self.client.post("/api/bakiye-siparis/delete-kpi", json=test_payload)
        self.assertEqual(del_admin.status_code, 200)
        self.assertEqual(del_admin.json()["durum"], "basarili")
        self.assertEqual(del_admin.json()["silinen_adet"], 1)

        # Silindiğini doğrula
        after_res = self.client.get(
            "/api/bakiye-siparis/kpi-records",
            params={"proje": "TEST-FILTRE-SIL-PROJESI", "sd": "99999999"}
        ).json()
        self.assertEqual(len(after_res["kayitlar"]), 0)

        self.client.delete("/api/identity")

    def test_bulk_template_and_analyze_and_tabs_ui(self):
        # 1) Ana sayfa HTML içinde Home, VennFlow, Toplu Veri Girişi sekmeleri ve KPI Pop-up modalı var mı?
        html_res = self.client.get("/")
        self.assertEqual(html_res.status_code, 200)
        html_text = html_res.text
        self.assertIn('id="pageTabBtnHome"', html_text)
        self.assertIn('id="pageTabBtnVennflow"', html_text)
        self.assertIn('id="pageTabBtnBulk"', html_text)
        self.assertIn('id="kpiRecordsModal"', html_text)

        # 2) Örnek toplu Excel şablonunu indir
        tpl_res = self.client.get("/api/bulk/template-excel")
        self.assertEqual(tpl_res.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            tpl_res.headers.get("content-type", "")
        )
        excel_bytes = tpl_res.content
        self.assertTrue(len(excel_bytes) > 1000)

        # 3) İndirilen örnek Excel'i geri yükleyip arka planda SD/Kalem eşleştirmesini ve Venn analizini doğrula
        analyze_res = self.client.post(
            "/api/bulk/analyze-excel",
            files={"file": ("VennFlow_Toplu_Veri_Girisi_Sablonu.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"engine_mode": "local", "threshold": "0.8"}
        )
        self.assertEqual(analyze_res.status_code, 200)
        adata = analyze_res.json()
        self.assertEqual(adata["durum"], "basarili")
        self.assertGreaterEqual(adata["toplam_grup_sayisi"], 1)

        first_grp = adata["gruplar"][0]
        self.assertEqual(first_grp["faaliyet_1_turu"], "Dry-Run")
        self.assertIn(first_grp["faaliyet_2_turu"], ("FKT", "SKT"))
        self.assertTrue(first_grp["sd"])
        self.assertTrue(first_grp["kalem"])
        self.assertTrue(first_grp["proje_adi"])
        self.assertTrue(first_grp["urun_kisa_metni"])
        self.assertIn("venn_verileri", first_grp)
        # Örnek şablonda 2 ortak, 2 yalnızca Dry-Run, 1 yalnızca FKT var
        self.assertEqual(first_grp["venn_verileri"]["intersection_count"], 2)

        # 4) Excel şablonunun A=SD, B=Kalem sütun yapısını ve E-G otomatik VLOOKUP formüllerini doğrula
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)
        ws = wb["Uygunsuzluk_Veri_Girisi"]
        self.assertIn("SD", str(ws["A1"].value))
        self.assertIn("Kalem", str(ws["B1"].value))
        self.assertIn("VLOOKUP(TRIM(A2)&\"|\"&TRIM(B2)", str(ws["E2"].value))
        self.assertIn("VLOOKUP(TRIM(A2)&\"|\"&TRIM(B2)", str(ws["F2"].value))
        self.assertIn("VLOOKUP(TRIM(A2)&\"|\"&TRIM(B2)", str(ws["G2"].value))

        # 5) VennFlow tekli önizlemede (/api/compare/excel-preview) SD ve Kalem sütunlarından bakiye bilgisi otomatik doluyor mu?
        preview_res = self.client.post(
            "/api/compare/excel-preview",
            files={"file": ("VennFlow_Toplu_Veri_Girisi_Sablonu.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        self.assertEqual(preview_res.status_code, 200)
        pdata = preview_res.json()
        self.assertIsNotNone(pdata.get("tespit_edilen_bakiye"))
        self.assertEqual(pdata["tespit_edilen_bakiye"]["sd"], first_grp["sd"])
        self.assertEqual(pdata["tespit_edilen_bakiye"]["kalem"], first_grp["kalem"])
        self.assertEqual(pdata["tespit_edilen_bakiye"]["proje_adi"], first_grp["proje_adi"])


if __name__ == "__main__":
    unittest.main()
